"""Read-only DOCX structure extractor used for proposal requirement audits."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}")
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip():
            print(f"P{index}: {paragraph.text}")
    for table_index, table in enumerate(document.tables):
        print(f"TABLE {table_index}")
        for row in table.rows:
            print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))


if __name__ == "__main__":
    main()
